"""Build the defence deck twice from one source.

Monarch_Defence_Deck.pptx         16 slides, no speaker notes.
Monarch_Defence_Deck_Script.pptx  the same 16 slides; each slide's speaker
                                  notes hold the script, every number on the
                                  slide with its meaning, and likely questions.
"""

import math as pymath
import os

from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from deck_theme import (BLUE, EMBER, GOLD, GREEN, HERE, LINE, M, MIDBLUE,
                        MUTED, TEXT, W, arrow, bar, card,
                        footer, header, icon, lattice, math, new_deck,
                        para, plain, rich, scale_bars, slide, textbox)

SECTIONS = ["The idea", "The method", "The results", "The physics",
            "What's next"]
DELTA_X = 0.1241
INSTRUMENT_URL = "https://monarch-4iy.pages.dev"
CORPUS_URL = INSTRUMENT_URL + "/corpus"
SCREENSHOT = os.path.join(HERE, "assets", "instrument_article.png")

GROUP_MEANS = [("Fear-activating", -0.0018, EMBER),
               ("Reward hook", -0.0128, MIDBLUE),
               ("High outrage", -0.0185, GOLD),
               ("Neutral", -0.0191, MUTED)]

SLIDES = [
    ("Title", None, 0.5,
     "Good morning. My name is Brian Mwai, and my dissertation is called "
     "Measuring the External Field. In simple words: when you read a news "
     "article, it pushes a little on what you think. My question is, can we "
     "measure how hard that push is, using only the article itself? Let me "
     "start with the idea behind it.",
     [], []),
    ("The idea", "The idea", 1,
     "Picture a room full of people, each for or against something. Physics "
     "has a model for this. It was built for magnets, where tiny arrows line "
     "up with their neighbours; here each arrow is a person. m is the mood of "
     "the room: minus one means everyone is against, plus one means everyone "
     "is for, zero means evenly split. J is how much people copy those around "
     "them. Beta is how seriously they take that pressure instead of deciding "
     "at random. And h is the outside push, the media, pushing everyone the "
     "same way at once. In every study I found, researchers guess h. Nobody "
     "measures it. So I asked: can we measure h from the article itself? When "
     "I tested my tools, five parts of my plan had to change.",
     ["m = the mood of the room: -1 all against, 0 split, +1 all for",
      "h = the media's push; 'field' here is general, not magnetic or "
      "electric"],
     [16]),
    ("Plan vs reality", "The idea", 1,
     "One: I planned to look at the amygdala, a fear centre deep in the "
     "brain, but the model only sees the brain's outer surface. Two: I "
     "planned to divide one brain area by another, but that broke on 69 of "
     "the 400 articles, a bit like dividing by zero, so I subtracted instead, "
     "which works on all 400. Three: I expected outrage to stand out, and fear "
     "did instead. Four: I planned to measure how strong the media's push is; "
     "the data couldn't, so I worked out the smallest it could be. Five: "
     "plain word counting sorted the articles better than my score, 98 "
     "against 63 out of 100, so my score is a measuring tool, not a fake-news "
     "detector. My supervisor approved all five. Now, how the measuring "
     "works.",
     ["69 of 400: articles where dividing failed (say 69, not 49)",
      "98 vs 63 out of 100: how often word counting vs my score ranks a "
      "manipulative article above a neutral one; 50 = coin toss"],
     [7, 17]),
    ("Tools and process", "The method", 1.5,
     "Here is how one article becomes one number. A computer voice reads it "
     "aloud, because the brain model learned from people watching and "
     "listening. A tool notes when each word is spoken. Then a brain model "
     "from Meta, called TRIBE, predicts how a typical adult brain would react "
     "at 20,484 spots on the brain's surface. From a standard brain map I take "
     "1,030 spots linked to emotion and 851 linked to careful thinking. My "
     "score is simple: the emotion reaction minus the thinking reaction. "
     "Above zero, emotion reacts more. Below zero, thinking reacts more. "
     "Nobody was put in a scanner; these are predictions. This article gets "
     "0.109 for emotion and 0.055 for thinking, so its score is plus 0.054, "
     "the highest of all 400. So which 400 articles?",
     ["20,484 = spots on the brain's surface the model predicts",
      "1,030 emotion spots, 851 thinking spots",
      "0.109 - 0.055 = +0.054: emotion reacts a bit more than thinking",
      "No unit: like marks on a ruler without cm; only comparisons matter"],
     [3, 6, 12]),
    ("The articles", "The method", 1,
     "I took 400 articles from four public collections, 100 each: fake news "
     "written to scare, angry political news, clickbait like 'you won't "
     "believe what happened next', and calm neutral writing, medical "
     "summaries and real news. All four groups average between 162 and 167 "
     "words, so no group wins just by being longer. Before starting, I worked "
     "out that 400 articles is enough to spot a difference as small as 0.027. "
     "And I ran every article twice, to check the results repeat. All 400 "
     "are online, so here is how to read that page.",
     ["162 to 167 = average words per group: same length",
      "0.027 = the smallest difference 400 articles can reliably spot "
      "(2.7% explained by group)",
      "x 2 = every article run twice"],
     [11, 13]),
    ("Reading the corpus", "The method", 1.5,
     "On the corpus page, every article is a row. Emotional is how strongly "
     "the emotion areas react. Deliberate is how strongly the thinking areas "
     "react. Score is emotional minus deliberate. Pre-scan label is what the "
     "original collection already called the article, before I touched it. "
     "Scores run from minus 0.070 to plus 0.054, so the gap from lowest to "
     "highest is 0.124. Remember that number; it comes back in the physics. "
     "On the right are the group averages. All four are below zero, so for "
     "every kind of article the thinking areas react more than the emotion "
     "areas. What differs is how far below zero. Fear is closest to zero, and "
     "neutral is furthest. Let me show you the live page. [SHOW CORPUS] So, "
     "are these differences real, or just chance?",
     ["-0.070 to +0.054 = lowest and highest score of the 400",
      "0.124 = gap from lowest to highest; used in the physics",
      "Group averages: fear -0.002, clickbait -0.013, outrage -0.018, "
      "neutral -0.019",
      "Below zero = thinking areas react more than emotion areas"],
     [18]),
    ("Result 1", "The results", 1,
     "To check, I used a standard test that compares groups, called a "
     "one-way ANOVA. It answers two plain questions. First: how much of the "
     "difference between articles comes from which group they are in? 10.7 "
     "percent. So the group matters, but about 89 percent comes from each "
     "individual article. Second: could this be pure chance? If the four "
     "groups were really the same, a result this strong would turn up about "
     "once in a billion tries. So chance is very unlikely. And running "
     "everything a second time gave 8.9 percent, close to the first. So the "
     "groups really differ. Which group makes the difference?",
     ["0.107 = 10.7% of the differences between articles come from their "
      "group; 89.3% from the article itself",
      "p = 1.0 x 10^-9 = if the groups were the same, this result would "
      "happen about once in a billion tries",
      "0.089 = the second run: 8.9%",
      "F = 15.779 = gap between groups is about 16 times the gap you'd "
      "expect from chance"],
     [14, 15]),
    ("Result 2", "The results", 1,
     "Fear does. Here I measure how far each group sits from the neutral "
     "articles, counted in typical spreads. The usual guide marks are 0.2, "
     "0.5 and 0.8. Fear sits 0.94 away, past the highest mark. Clickbait sits "
     "0.32 away. Outrage only 0.03, almost on top of neutral. Why is outrage "
     "so close? Look at the two small charts. Fear switches on the emotion "
     "areas, plus 0.61, and not the thinking areas, minus 0.10, so the gap "
     "grows and the score moves. Outrage switches on both, plus 0.51 and plus "
     "0.49. Both rise together, so the gap, and the score, stay flat. Outrage "
     "does affect the brain, just on both sides equally. Next: can we trust "
     "these scores?",
     ["d = distance from neutral, in typical spreads (standard deviations)",
      "Fear 0.94, clickbait 0.32, outrage 0.03; guide marks 0.2, 0.5, 0.8",
      "Fear: emotion +0.61, thinking -0.10, so the gap grows",
      "Outrage: emotion +0.51, thinking +0.49, so the gap stays near 0"],
     []),
    ("Result 3", "The results", 1,
     "I ran all 400 articles twice. The two runs agree at 0.87, on a scale "
     "where 1 means identical and 0 means no agreement at all. But 51 "
     "articles flipped: positive one time, negative the other. Is that a "
     "problem? If only small random noise were at work, we'd expect about 55 "
     "flips, anywhere from 44 to 67. We saw 51, inside that range. And the "
     "flipped articles all sit very close to zero, like a coin balanced on "
     "its edge, where a tiny wobble tips it. So I trust group averages, and I "
     "never judge a single article. Now the most important honest check.",
     ["0.8725 = agreement between the two runs; 1 = identical, 0 = none",
      "51 of 400 = 1 in 8 articles flipped sign",
      "55 (44 to 67) = flips that random noise alone would cause",
      "Flipped articles average 0.003 from zero; the rest 0.022"],
     []),
    ("Honest check", "The results", 1,
     "Can my score tell manipulative articles from neutral ones? The test is "
     "simple: pick one manipulative and one neutral article at random. How "
     "often does the method rank the manipulative one higher? 50 out of 100 "
     "is a coin toss, and 100 is perfect. A standard sentiment tool gets 54. "
     "My score gets 63. Plain word counting gets 98. Why is word counting so "
     "good? Each group came from a different source, and the words give the "
     "source away: from the words alone you can tell the source 66 times out "
     "of 100, where guessing gets 25. So the difference between groups can't "
     "be put down to writing style alone, because the source is mixed in. "
     "That's why my score is a measurement for the physics, not a fake-news "
     "detector. Next: does the brain model match real brains?",
     ["AUC = out of 100 random pairs, how often the manipulative article "
      "ranks higher",
      "54 sentiment tool, 63 my score, 98 word counting; 50 = coin toss",
      "66% = words reveal the source; guessing = 25%"],
     [5]),
    ("Real brains", "The results", 1,
     "I tested the model on real brain scans of four people watching the TV "
     "show Friends. r measures how closely two signals move together: 0 is no "
     "match, 1 is perfect. Even real people don't match each other perfectly: "
     "over the full episode they match at 0.152, so that's the best any model "
     "could hope for. On the 2-minute clip I tested, the best possible is "
     "0.096, and the model reaches 0.028, about 30 percent of the best "
     "possible. The p-value is 0.048, just under the usual 5 percent line. "
     "And it's positive: the model moves in the same direction as real "
     "brains, not the opposite, as an outside audit had claimed. I also found "
     "and fixed three bugs along the way. The full-episode test is next. Now "
     "the physics.",
     ["r = how closely two signals move together; 0 none, 1 perfect",
      "0.152 = real people vs each other: the best possible",
      "0.096 best possible on the clip; the model 0.028 = 30% of it",
      "p = 0.048 = about 5 in 100 chance of a match this good by luck alone"],
     []),
    ("The minimum push", "The physics", 1.5,
     "Back to the room of people. The media's push, h, equals my score times "
     "a number called alpha: how much push you get per unit of score. I "
     "can't measure alpha, so I asked a different question: how big would "
     "alpha have to be for the media to flip the majority of the room? That "
     "depends on how much people copy each other, beta J. The more they copy "
     "each other, the more stubborn the room, and the bigger the push needed. "
     "At beta J of 2, the push needed is 0.533. My scores only span 0.124. So "
     "alpha must be at least 0.533 divided by 0.124, which is 4.29. In plain "
     "words: anyone claiming these articles flip opinion must assume a media "
     "strength of at least 4.29, or the model says it can't happen. It's a "
     "minimum bar, not a claim that media does flip opinion. So what could "
     "this be used for?",
     ["alpha = how much push per unit of score",
      "beta J = 2: how much people copy each other; say 'beta J', not 'beta'",
      "0.533 = push needed to flip the room at beta J = 2",
      "0.533 / 0.124 = 4.29 = the minimum alpha"],
     [9, 10, 16]),
    ("Children's content", "What's next", 1,
     "When I presented my proposal, I was asked: could this analyse "
     "anything, even children's TV? The tool takes video and audio, so it "
     "could score a cartoon. But the brain model learned from adults, so it "
     "can't tell us how a child reacts. I only tested news text. And about 1 "
     "in 8 single scores flip between runs, so one show on its own isn't "
     "reliable. For children we'd first need children's brain data, with "
     "ethics approval. Each limit like this points to a next study.",
     ["1 in 8 = 51 of 400 single scores flip between runs"],
     [1, 2, 8, 19]),
    ("Limits", "What's next", 0.5,
     "Five limits, each with its fix. These are predictions, so check them "
     "against more real brain scans. Group and source are mixed, so build a "
     "new set where every source has every type. The model misses deep brain "
     "areas, so use a model that sees them. Alpha isn't measured, so compare "
     "scores with real opinion polls. And the brain test was 2 minutes, so "
     "run the full episode. That leads to where the work goes next.",
     ["5 limits, each stated in the thesis"], []),
    ("Taking it further", "What's next", 0.75,
     "Three directions. Publish: three papers, shared publicly first. "
     "Extend: a Kenyan case, news coverage of the 2024 Finance Bill, plus "
     "video and audio, and measuring alpha against opinion polls. And "
     "opportunities: further study, work with brain and media researchers, "
     "and a tool that is already online. To close, four conclusions.",
     ["Kenya case needs GPU funding"],
     [4, 13]),
    ("Conclusions", "What's next", 1,
     "Four things to take away. One: you can measure a push from an "
     "article's content, and it repeats: the two runs agree at 0.87 out of "
     "1. Two: it tells the four kinds of articles apart; the group explains "
     "10.7 percent of the differences, with the source mixed in. Three: we "
     "can't yet measure media's strength, but we know its minimum: 4.29 when "
     "people copy each other at strength 2. Four: the brain model moves with "
     "real brains, at 30 percent of the best possible. Three papers come from "
     "this work. Thank you, and I'm happy to take questions.",
     ["0.87 = the two runs agree", "10.7% = explained by group",
      "4.29 = minimum media strength at beta J = 2",
      "30% = the model's match vs the best possible"],
     [3, 5, 19]),
]

QUESTIONS = [
    ("Can this analyse any content, even what children watch?",
     "It can score video and audio. But the model learned from adult brains, "
     "so it can't say how a child reacts, and I only tested news. It would "
     "need children's brain data and ethics approval first."),
    ("Could a parent or a regulator use it to rate one show?",
     "Not yet. About 1 in 8 single scores flip between runs, so one score "
     "isn't reliable. It is for comparing groups of content."),
    ("Was anyone scanned? Is it reading minds?",
     "No. Nobody was scanned. The scores are predictions of how a typical "
     "adult brain reacts. It rates the content, never a person."),
    ("Could someone use it to make content more manipulative?",
     "That's a real risk for any measure like this. That's why it only "
     "reports group results, for research, and the model's licence forbids "
     "commercial use."),
    ("If word counting does better, why not use that?",
     "Word counting mostly learns where an article came from: it names the "
     "source 66 times out of 100. My score is meant to be the media push in "
     "the physics model, not a detector. I include the comparison to be "
     "honest."),
    ("Why read the articles aloud?",
     "The brain model learned from people watching and listening, so it "
     "needs speech and the timing of each word."),
    ("Why no amygdala?",
     "The model only predicts the brain's outer surface. The amygdala is "
     "deep inside, so I make no claim about it."),
    ("What would it take for children's media?",
     "Children's brain scans with ethics approval and parental consent, a "
     "set of children's videos, and content where source and type are "
     "mixed."),
    ("Why is this physics?",
     "The opinion model is the same maths physics uses for magnets lining "
     "up. My part is measuring the push and working out its minimum size."),
    ("What does alpha >= 4.29 mean in practice?",
     "It's a minimum bar. Anyone who claims these articles flip a group's "
     "opinion, where people copy each other at strength 2, must use a media "
     "strength of at least 4.29, or the model says it can't happen."),
    ("Why only 400 articles?",
     "I worked out beforehand that 400 is enough to spot a difference as "
     "small as 2.7 percent. Each article takes about 70 seconds on the "
     "computer, and I ran them all twice."),
    ("Is it okay to use Meta's model?",
     "Yes, for research. Its licence allows non-commercial use, I cite it "
     "throughout, and this project makes no money."),
    ("Why didn't you use Kenyan media?",
     "Cost. The 400 articles, run twice, used about 16 hours of free "
     "computer time. Kenyan articles would also have to be labelled by hand, "
     "by paid people, because no labelled set exists. It is the first "
     "extension once there is funding."),
    ("What is eta squared, physically?",
     "It isn't a physical quantity. It's a share between 0 and 1: how much "
     "of the total difference between articles is explained by their group. "
     "0.107 means 10.7 percent."),
    ("Why a one-way ANOVA?",
     "Because there is one thing sorting the articles, the group, and one "
     "number per article, the score. That's exactly what a one-way ANOVA "
     "tests. I also compared each group with neutral separately."),
    ("What does beta J mean, and why must a be negative?",
     "Beta J is how strongly people copy each other. When it's above 1, the "
     "room settles into two stable opinions, for or against; in the maths "
     "that shows up as a negative coefficient, a. Then a push is needed to "
     "flip from one to the other. Alpha itself is positive."),
    ("Why a difference, not a ratio?",
     "The model's numbers are often near zero or negative, so dividing by "
     "them breaks, like dividing by zero. It failed on 69 of 400 articles. "
     "Subtracting works for all 400."),
    ("Is the difference between groups just writing style?",
     "It can't be put down to style alone, because each group came from one "
     "source, and the words reveal the source 66 times out of 100. A new set "
     "where every source has every type would separate the two."),
    ("What about ethics?",
     "I collected no data from any person. The brain model was built by "
     "Meta from public scans of volunteers who consented, and I used only "
     "public articles and public scans. The tool rates content, never a "
     "viewer."),
]

EQUATIONS = {
    1: [("Ising energy", "E = -J sum(s_i s_j) - h sum(s_i). Each person s_i is +1 or "
         "-1; agreeing with neighbours lowers the energy, and h rewards "
         "pointing the media's way"),
        ("Mean-field equation", "m = tanh(beta J m + h). Each person feels "
         "the room's average m instead of each neighbour; solving it gives "
         "the room's mood")],
    3: [("The score", "X = A_emotional - A_deliberate, each A the average "
         "over its spots. A difference, not a ratio, because A is often near "
         "zero or negative")],
    4: [("Power analysis", "the smallest effect a study can catch: with 4 "
         "groups of 100, p < 0.05 and 80% power, that is eta squared = "
         "0.027")],
    5: [("Spread", "Delta X = X_max - X_min = 0.054 - (-0.070) = 0.124. It "
         "is the biggest push difference the articles can give")],
    6: [("Eta squared", "eta^2 = SS_between / SS_total: the share of all "
         "variation explained by group"),
        ("F-test", "F = (between-group variance) / (within-group variance) "
         "= 15.779; p is the chance of an F this big if the groups were the "
         "same")],
    7: [("Cohen's d", "d = (group mean - neutral mean) / pooled standard "
         "deviation: the gap in units of typical spread")],
    8: [("ICC", "ICC = (variance between articles) / (between + run-to-run "
         "variance) = 0.8725"),
        ("Noise model", "run-to-run noise sigma = sd(run1 - run2) / "
         "sqrt(2) = 0.0073; simulating that noise gives 55 expected flips")],
    9: [("AUC", "AUC = P(X_manipulative > X_neutral) for a random pair; "
         "0.5 is chance")],
    10: [("Pearson r", "r = cov(prediction, brain) / (sd_prediction x "
          "sd_brain): how closely they move together"),
          ("Noise ceiling", "how well one person's brain predicts the "
           "others'; no model can beat it")],
    11: [("Landau free energy", "F(m) = a m^2 + b m^4 - h m, with a = (1 - "
          "beta J)/2 and b = 1/12. When beta J > 1, a < 0 and two stable "
          "opinions (two valleys) appear"),
         ("Critical field", "from m = tanh(beta J m + h), h = atanh(m) - "
          "beta J m. The valley vanishes where dh/dm = 0, at m*^2 = 1 - 1/beta "
          "J; h_c = |atanh(m*) - beta J m*| = 0.533 at beta J = 2"),
         ("The bound", "the strongest push the articles give is h = alpha x "
          "Delta X. To flip the room it must reach h_c, so alpha >= h_c / "
          "Delta X = 0.533 / 0.124 = 4.29")],
}

KEY_POINTS = [
    "An article pushes on what people think. Physics guesses that push; this "
    "work measures it and finds its minimum size.",
    "Physics models opinion like magnets. The media's push, h, is always "
    "guessed. Can we measure it?",
    "Testing the tools changed five parts of the plan, all approved.",
    "Score = emotion reaction minus thinking reaction, predicted by a brain "
    "model. Nobody was scanned.",
    "400 articles, 4 groups of 100, same length, each run twice.",
    "Every column explained. All four groups sit below zero: thinking reacts "
    "more than emotion.",
    "The groups really differ: group explains 10.7%, and chance is about 1 "
    "in a billion.",
    "Fear moves the score (0.94). Outrage raises emotion and thinking "
    "equally, so its score stays flat (0.03).",
    "Two runs agree at 0.87 out of 1. The 51 flips are what noise predicts, "
    "so judge groups, never one article.",
    "Word counting wins, 98 vs 63 out of 100, because source is mixed in. "
    "The score measures, it doesn't detect.",
    "The brain model moves with real brains: 30% of the best possible.",
    "Media would need a strength of at least 4.29 to flip the room.",
    "It could score children's TV, but it isn't valid for children yet.",
    "Every limit has a fix.",
    "Publish, extend to Kenya and video, and further study.",
    "Four take-aways, then questions.",
]

SO_WHAT = (
    "Opinion models treat media as a push, but everyone guesses how big it "
    "is. I built a way to measure a candidate for that push from the content "
    "itself. It separates the four groups at 10.7%, with source mixed in, so "
    "I don't claim it detects manipulation. And even without knowing the "
    "exact strength, the measured spread gives a floor: media would need a "
    "coupling of at least 4.29 to flip a group with beta J = 2. That turns a "
    "guessed number into a testable one.")
SO_WHAT_SLIDES = {0, 11, 15}

DEFINITIONS = {
    0: [("Sociophysics", "using the tools of statistical physics to model how "
                         "groups of people behave")],
    1: [("Ising model", "a model of magnets where each spin points up or "
                        "down and prefers to match its neighbours; used here "
                        "for two-sided opinions"),
        ("Mean field", "each person feels the average of the whole group, "
                       "not particular neighbours; it makes the maths "
                       "solvable"),
        ("tanh", "a smooth S-shaped function that keeps m between -1 and "
                 "+1")],
    2: [("Amygdala", "a small structure deep in the brain linked to fear"),
        ("Cortex", "the brain's outer surface, where the model predicts"),
        ("Amendment", "the approved change to the proposal's scope and "
                      "title")],
    3: [("TRIBE v2", "Meta's published brain model: it predicts fMRI "
                     "responses from text, audio and video"),
        ("fMRI", "brain scanning that tracks blood flow as a stand-in for "
                 "activity"),
        ("fsaverage5", "a standard brain surface with 20,484 points"),
        ("Glasser 2016 atlas", "a standard map dividing the cortex into "
                               "named areas"),
        ("Standardised activity", "the model's output, scaled so it has no "
                                  "physical unit")],
    4: [("Power analysis", "working out in advance how big a study must be "
                           "to detect an effect of a given size")],
    5: [("Pre-scan label", "the label an article carried in its source "
                           "collection before anything was scanned")],
    6: [("One-way ANOVA", "compares the averages of several groups defined "
                          "by one factor"),
        ("F", "spread between group averages divided by spread within "
              "groups"),
        ("eta squared", "between-group sum of squares over total sum of "
                        "squares; a share from 0 to 1"),
        ("p-value", "the probability of a result at least this large if "
                    "there were no real difference")],
    7: [("Cohen's d", "difference between two group averages divided by the "
                      "typical spread (standard deviation)")],
    8: [("ICC", "intraclass correlation: agreement between repeated "
                "measurements of the same items; 1 is identical"),
        ("Sign reversal", "a score positive on one run and negative on the "
                          "other")],
    9: [("AUC", "area under the ROC curve: chance a random manipulative "
                "article outscores a random neutral one"),
        ("TF-IDF", "word counting that weights words by how distinctive "
                   "they are"),
        ("VADER", "a standard sentiment tool"),
        ("Confound", "a second factor, here source, that moves with the one "
                     "studied")],
    10: [("Noise ceiling", "how well real brains predict each other; the "
                           "best any model can do"),
         ("r", "Pearson correlation: how closely two signals rise and fall "
               "together"),
         ("Circular-shift test", "shift one signal in time 61 times to see "
                                 "how often chance gives a match this good")],
    11: [("Critical field", "the smallest field that flips the majority"),
         ("Lower bound", "a minimum: alpha can't be smaller than this")],
    12: [("Ethics approval", "formal permission from a review board, "
                             "required for research with children")],
    13: [("Crossed design", "every source contributes every type of "
                            "article")],
    14: [("Preprint", "a paper shared publicly before peer review, e.g. on "
                      "arXiv")],
}


def notes_for(index):
    title, _, minutes, script, cues, questions = SLIDES[index]
    lines = ["KEY POINT: " + KEY_POINTS[index], "",
             "SCRIPT (about %g min)" % minutes, script]
    if cues:
        lines += ["", "WHAT THE NUMBERS MEAN"]
        lines += ["- " + cue for cue in cues]
    if index in EQUATIONS:
        lines += ["", "EQUATION SIDE NOTE"]
        lines += ["- %s: %s." % pair for pair in EQUATIONS[index]]
    if index in DEFINITIONS:
        lines += ["", "DEFINITIONS (if asked)"]
        lines += ["- %s: %s." % pair for pair in DEFINITIONS[index]]
    if index in SO_WHAT_SLIDES:
        lines += ["", "IF ASKED \"SO WHAT?\"", SO_WHAT]
    if questions:
        lines += ["", "IF ASKED"]
        for q in questions:
            question, answer = QUESTIONS[q - 1]
            lines += ["Q: " + question, "A: " + answer, ""]
    return "\n".join(lines).rstrip()


def banner(s, y, lead, text, colour=GOLD, h=0.62, size=16):
    shp = card(s, M, y, W - 2 * M, h, line=colour)
    tf = shp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    rich(tf, [(lead + "  ", True, colour), (text, False, TEXT)], size=size,
         first=True, space_after=0)


def chip(s, x, y, w, h, text, colour=TEXT, line=LINE, size=15, bold=False,
         align=PP_ALIGN.LEFT, icon_name=None):
    shp = card(s, x, y, w, h, line=line)
    tf = shp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if icon_name:
        icon(s, icon_name, x + 0.18, y + (h - 0.3) / 2, h=0.3)
        tf.margin_left = Inches(0.6)
    para(tf, text, size=size, bold=bold, colour=colour, first=True,
         align=align, space_after=0)
    return shp


def stat(s, x, y, w, h, value, label, colour=BLUE, size=36, line=None,
         label_size=13):
    shp = card(s, x, y, w, h, line=line or colour)
    tf = shp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, value, size=size, bold=True, colour=colour, first=True,
         align=PP_ALIGN.CENTER, space_after=0)
    para(tf, label, size=label_size, colour=TEXT, align=PP_ALIGN.CENTER,
         space_after=0)


def style_chart_text(chart, size=12):
    chart.font.size = Pt(size)
    chart.font.color.rgb = MUTED
    chart.font.name = "Calibri"


def donut(s, x, y, d, share, colour):
    data = CategoryChartData()
    data.categories = ["group", "rest"]
    data.add_series("share", (share, 1 - share))
    frame = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y),
                               Inches(d), Inches(d), data)
    chart = frame.chart
    chart.has_legend = False
    chart.has_title = False
    plot = chart.plots[0]
    for pt, fill in zip(plot.series[0].points, (colour, LINE)):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = fill
        pt.format.line.fill.background()
    hole = plot._element.find(qn("c:holeSize"))
    if hole is None:
        hole = OxmlElement("c:holeSize")
        plot._element.append(hole)
    hole.set("val", "70")


def critical_field(k):
    if k <= 1:
        return 0.0
    m = pymath.sqrt(1 - 1 / k)
    return abs(pymath.atanh(m) - k * m)


def alpha_curve(s, x, y, w, h):
    data = XyChartData()
    curve = data.add_series("alpha required")
    k = 1.0
    while k <= 2.5001:
        curve.add_data_point(round(k, 3), critical_field(k) / DELTA_X)
        k += 0.05
    point = data.add_series("our case")
    point.add_data_point(2.0, critical_field(2.0) / DELTA_X)
    frame = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_SMOOTH_NO_MARKERS,
                               Inches(x), Inches(y), Inches(w), Inches(h), data)
    chart = frame.chart
    chart.has_legend = False
    style_chart_text(chart)
    line_series, dot_series = chart.plots[0].series
    line_series.format.line.color.rgb = BLUE
    line_series.format.line.width = Pt(3)
    line_series.smooth = True
    dot_series.format.line.fill.background()
    dot_series.marker.style = XL_MARKER_STYLE.CIRCLE
    dot_series.marker.size = 12
    dot_series.marker.format.fill.solid()
    dot_series.marker.format.fill.fore_color.rgb = EMBER
    dot_series.marker.format.line.fill.background()
    for axis, lo, hi, unit in ((chart.category_axis, 1.0, 2.5, 0.5),
                               (chart.value_axis, 0, 8, 2)):
        axis.minimum_scale, axis.maximum_scale = lo, hi
        axis.major_unit = unit
        axis.format.line.color.rgb = LINE
        axis.has_major_gridlines = axis is chart.value_axis
    chart.value_axis.major_gridlines.format.line.color.rgb = LINE


def next_line(s, text):
    tf = textbox(s, W - M - 6.0, 6.62, 6.0, 0.3, align=PP_ALIGN.RIGHT)
    para(tf, "Next: " + text + "  →", size=12, bold=True, colour=GOLD,
         first=True, align=PP_ALIGN.RIGHT, space_after=0)


def build_deck(with_notes):
    prs = new_deck()

    def page(i):
        return slide(prs, notes=notes_for(i) if with_notes else None)

    def done(s, i):
        footer(s, SECTIONS, SLIDES[i][1], i + 1)

    # 1 title
    s = page(0)
    lattice(s, 8.75, 1.15, 7, 9, 0.52, seed=7, bias=0.64)
    tf = textbox(s, 8.75, 5.95, 3.64, 0.5, align=PP_ALIGN.CENTER)
    para(tf, "each arrow is one person's opinion", size=13, colour=MUTED,
         first=True, align=PP_ALIGN.CENTER, space_after=0)
    tf = textbox(s, M + 0.15, 1.15, 7.6, 0.4)
    para(tf, "B.Sc. PHYSICS   ·   CUEA   ·   DISSERTATION DEFENCE", size=13,
         bold=True, colour=BLUE, first=True, space_after=0)
    tf = textbox(s, M + 0.15, 1.65, 7.6, 1.9)
    para(tf, "Measuring the", size=50, bold=True, colour=TEXT, first=True,
         space_after=0)
    para(tf, "External Field", size=50, bold=True, colour=BLUE, space_after=0)
    tf = textbox(s, M + 0.15, 3.55, 7.4, 1.0)
    para(tf, "Can we measure how hard a news article pushes on what people "
             "think?", size=21, colour=MUTED, first=True, space_after=0)
    bar(s, M + 0.15, 4.75, 1.4, 0.05, EMBER)
    tf = textbox(s, M + 0.15, 5.0, 7.4, 1.4)
    rich(tf, [("Brian Mwai", True, TEXT), ("   1050555", False, MUTED)],
         size=20, first=True, space_after=4)
    para(tf, "Supervisor: Dr. Songa Mutambi", size=15, colour=MUTED,
         space_after=2)
    para(tf, "25 September 2026", size=15, colour=MUTED, space_after=0)

    # 2 the idea
    s = page(1)
    header(s, "The idea", "Physics treats media as a push, h, on opinion. "
                          "Nobody measures h", "waveform")
    math(s, r"m = \tanh(\beta J\, m + h)", M, 1.65, size=36)
    for i, (sym, text, colour) in enumerate([
            ("m", "magnetisation\n= average opinion, −1 to +1", BLUE),
            ("J", "coupling\n= pull toward neighbours", BLUE),
            (r"\beta", "inverse temperature\n= how firmly people follow",
             BLUE),
            ("h", "external field\n= media's push on everyone", GOLD)]):
        x = M + (i % 2) * 3.85
        y = 2.75 + (i // 2) * 1.12
        card(s, x, y, 3.7, 0.98, line=GOLD if sym == "h" else LINE)
        math(s, sym, x + 0.1, y + 0.2, size=30, colour=colour, centre_w=0.8)
        tf = textbox(s, x + 0.95, y + 0.12, 2.7, 0.8, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, text, size=14, colour=TEXT, first=True, space_after=0)
    lattice(s, 8.9, 1.75, 6, 6, 0.55, seed=3, bias=0.66)
    plain(s.shapes.add_shape(MSO_SHAPE.UP_ARROW, Inches(12.4), Inches(2.3),
                             Inches(0.4), Inches(2.2)), GOLD)
    math(s, "h", 12.4, 4.6, size=28, colour=GOLD, centre_w=0.4)
    banner(s, 5.2, "Question:", "can h be measured from the content itself?",
           colour=EMBER, h=0.75)
    next_line(s, "what testing the tools changed")
    done(s, 1)

    # 3 plan vs reality
    s = page(2)
    header(s, "The idea", "Testing the tools changed five assumptions",
           "magnifying-glass")
    for x, label, colour in [(M, "WE PLANNED", MUTED), (5.55, "WE FOUND", BLUE)]:
        tf = textbox(s, x, 1.45, 4, 0.3)
        para(tf, label, size=11, bold=True, colour=colour, first=True,
             space_after=0)
    for i, (planned, found) in enumerate([
            ("Measure the amygdala", "Model predicts the brain surface only, "
                                     "so we used surface areas"),
            ("Score as a ratio", "Ratio undefined for 69 of 400, so we used "
                                 "a difference"),
            ("Outrage stands out most", "Fear stands out; outrage lifts both "
                                        "areas equally"),
            ("Measure α, media's strength", "Data can't pin α, so we derived "
                                            "its minimum"),
            ("A manipulation detector", "Word counting 0.98 vs our 0.63, so "
                                        "it's a measurement")]):
        y = 1.8 + i * 0.86
        chip(s, M, y, 4.1, 0.72, planned, colour=MUTED, icon_name="x-circle")
        arrow(s, 4.95, y + 0.29, 0.5, 0.16, EMBER)
        chip(s, 5.55, y, 7.03, 0.72, found, line=BLUE, icon_name="check-circle")
    tf = textbox(s, M, 6.15, 6.0, 0.35)
    para(tf, "All five are in the approved amendment.", size=13, colour=MUTED,
         first=True, space_after=0)
    next_line(s, "how the measurement works")
    done(s, 2)

    # 4 tools and process
    s = page(3)
    header(s, "The method", "How one article becomes one number", "brain")
    steps = [("newspaper", "Article", "400 news texts"),
             ("waveform", "Read aloud", "text-to-speech"),
             ("text-aa", "Word timings", "timing tool"),
             ("brain", "TRIBE v2", "20,484 surface points"),
             ("target", "Brain map", "1,030 + 851 points")]
    sw, gap = 2.12, 0.31
    for i, (icon_name, head, tool) in enumerate(steps):
        x = M + i * (sw + gap)
        card(s, x, 1.45, sw, 1.3, line=BLUE if i == 3 else LINE)
        icon(s, icon_name, x + sw / 2 - 0.2, 1.55, h=0.4)
        tf = textbox(s, x, 1.98, sw, 0.75, align=PP_ALIGN.CENTER)
        para(tf, head, size=15, bold=True, colour=TEXT, first=True,
             space_after=0)
        para(tf, tool, size=12, colour=MUTED, space_after=0)
        if i < len(steps) - 1:
            arrow(s, x + sw + 0.05, 2.02, gap - 0.1, 0.16)
    card(s, M, 2.95, 5.3, 3.15, line=BLUE)
    s.shapes.add_picture(SCREENSHOT, Inches(M + 0.1), Inches(3.03),
                         width=Inches(5.1))
    tf = textbox(s, M, 6.13, 5.4, 0.5)
    rich(tf, [("One article: ", True, TEXT),
              ("emotional +0.109 − deliberate +0.055 = score +0.054, "
               "highest of 400", False, MUTED)], size=11, first=True,
         space_after=0)
    card(s, 6.3, 2.95, 6.28, 1.95, line=BLUE)
    tf = textbox(s, 6.5, 3.03, 5, 0.3)
    para(tf, "THE SCORE, X", size=11, bold=True, colour=MUTED, first=True,
         space_after=0)
    math(s, r"X = \bar{A}_{\mathrm{emotional}} - \bar{A}_{\mathrm{deliberate}}",
         6.3, 3.3, size=22, centre_w=6.28)
    bar(s, 6.6, 4.12, 2.8, 0.1, BLUE)
    bar(s, 9.43, 4.12, 2.8, 0.1, EMBER)
    bar(s, 9.4, 3.99, 0.03, 0.36, TEXT)
    for x, text, colour, al in [(6.6, "X < 0: reasoning leads", BLUE,
                                 PP_ALIGN.LEFT),
                                (9.43, "X > 0: emotion leads", EMBER,
                                 PP_ALIGN.RIGHT)]:
        tf = textbox(s, x, 4.24, 2.8, 0.3, align=al)
        para(tf, text, size=12, bold=True, colour=colour, first=True,
             space_after=0)
    tf = textbox(s, 6.5, 4.52, 5.9, 0.3, align=PP_ALIGN.CENTER)
    para(tf, "Ā = average predicted activity; standardised, no unit  ·  "
             "nobody scanned", size=11, colour=MUTED, first=True,
         space_after=0)
    link_card = card(s, 6.3, 5.05, 6.28, 1.05, line=GOLD)
    link_card.click_action.hyperlink.address = INSTRUMENT_URL
    icon(s, "target", 6.5, 5.33, h=0.46)
    tf = textbox(s, 7.15, 5.1, 5.3, 1.0, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "TRY THE INSTRUMENT LIVE", size=11, bold=True, colour=GOLD,
         first=True, space_after=0)
    para(tf, INSTRUMENT_URL.replace("https://", ""), size=22, bold=True,
         colour=TEXT, space_after=0)
    para(tf, "code and data: github.com/brn-mwai/monarch", size=11,
         colour=MUTED, space_after=0)
    next_line(s, "the 400 articles")
    done(s, 3)

    # 5 the articles
    s = page(4)
    header(s, "The method", "400 articles, four groups, matched in length",
           "newspaper")
    for i, (name, src, colour) in enumerate([
            ("Fear-activating", "fake news (ISOT)", EMBER),
            ("High outrage", "partisan news (SemEval)", GOLD),
            ("Reward hook", "clickbait (Webis)", MIDBLUE),
            ("Neutral", "medical + real news", MUTED)]):
        x = M + i * (2.72 + 0.317)
        card(s, x, 1.55, 2.72, 1.9)
        bar(s, x, 1.55, 2.72, 0.07, colour)
        tf = textbox(s, x, 1.75, 2.72, 1.6, align=PP_ALIGN.CENTER)
        para(tf, "100", size=40, bold=True, colour=colour, first=True,
             space_after=0)
        para(tf, name, size=17, bold=True, colour=TEXT, space_after=0)
        para(tf, src, size=12, colour=MUTED, space_after=0)
    for i, (value, label, colour) in enumerate([
            ("162–167", "average words per article,\nlowest to highest group",
             BLUE),
            ("0.027", "smallest η² 400 articles can\ndetect, set before "
                      "scanning", GREEN),
            ("× 2", "every article\nscanned twice", GOLD)]):
        stat(s, M + i * (3.75 + 0.29), 3.8, 3.75, 1.75, value, label,
             colour=colour, size=40)
    tf = textbox(s, M, 5.8, W - 2 * M, 0.35)
    rich(tf, [("All 400, with scores:  ", True, TEXT),
              (CORPUS_URL.replace("https://", ""), True, GOLD)], size=14,
         first=True, space_after=0)
    next_line(s, "reading the corpus page")
    done(s, 4)

    # 6 reading the corpus
    s = page(5)
    header(s, "The method", "Reading the corpus page: what each number means",
           "newspaper")
    card(s, M, 1.45, 6.2, 4.6, line=BLUE)
    tf = textbox(s, M + 0.25, 1.55, 5.8, 0.3)
    para(tf, "COLUMN   →   MEANING", size=11, bold=True, colour=MUTED,
         first=True, space_after=0)
    for i, (column, meaning) in enumerate([
            ("Emotional", "average predicted activity, 1,030 emotional "
                          "points"),
            ("Deliberate", "average predicted activity, 851 reasoning "
                           "points"),
            ("Score", "Emotional − Deliberate = X"),
            ("Pre-scan label", "label from its source, given before any "
                               "scan"),
            ("Words", "article length")]):
        y = 1.95 + i * 0.66
        chip(s, M + 0.2, y, 1.9, 0.54, column, colour=GOLD, size=14,
             bold=True, align=PP_ALIGN.CENTER)
        tf = textbox(s, M + 2.25, y, 3.8, 0.54, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, meaning, size=13, colour=TEXT, first=True, space_after=0)
    tf = textbox(s, M + 0.25, 5.3, 5.8, 0.7)
    rich(tf, [("All 400 scores: ", True, TEXT),
              ("−0.070 to +0.054, spread ΔX = 0.124", False, TEXT)], size=13,
         first=True, space_after=2)
    para(tf, "Units: standardised activity, no physical unit", size=11,
         colour=MUTED, space_after=0)
    card(s, 7.2, 1.45, 5.38, 4.6, line=LINE)
    tf = textbox(s, 7.4, 1.55, 5.0, 0.3)
    para(tf, "AVERAGE SCORE X PER GROUP", size=11, bold=True, colour=MUTED,
         first=True, space_after=0)
    zero, scale = 12.1, 0.02 / 3.2
    bar(s, zero, 1.95, 0.02, 3.1, TEXT)
    tf = textbox(s, zero - 0.5, 5.05, 1.0, 0.3, align=PP_ALIGN.CENTER)
    para(tf, "0", size=12, bold=True, colour=TEXT, first=True,
         align=PP_ALIGN.CENTER, space_after=0)
    for i, (name, mean, colour) in enumerate(GROUP_MEANS):
        y = 2.05 + i * 0.75
        length = abs(mean) / scale
        bar(s, zero - length, y + 0.3, length, 0.3, colour)
        tf = textbox(s, 7.4, y - 0.02, 3.2, 0.3)
        para(tf, name, size=13, bold=True, colour=TEXT, first=True,
             space_after=0)
        tf = textbox(s, zero - length - 1.25, y + 0.26, 1.2, 0.35,
                     align=PP_ALIGN.RIGHT)
        para(tf, "%+.3f" % mean, size=13, bold=True, colour=colour,
             first=True, align=PP_ALIGN.RIGHT, space_after=0)
    tf = textbox(s, 7.4, 5.35, 5.0, 0.6)
    para(tf, "All four below 0: the deliberate areas respond more in every "
             "group.", size=12, colour=MUTED, first=True, space_after=0)
    tf = textbox(s, M, 6.2, 6.5, 0.35)
    rich(tf, [("Live: ", True, TEXT),
              (CORPUS_URL.replace("https://", ""), True, GOLD)], size=13,
         first=True, space_after=0)
    next_line(s, "are these differences real?")
    done(s, 5)

    # 7 result 1
    s = page(6)
    header(s, "The results", "Result 1: the four groups differ, η² = 0.107",
           "chart-bar")
    donut(s, M, 1.45, 3.7, 0.1068, BLUE)
    tf = textbox(s, M, 2.8, 3.7, 1.0, align=PP_ALIGN.CENTER)
    para(tf, "10.7%", size=40, bold=True, colour=TEXT, first=True,
         space_after=0)
    tf = textbox(s, M, 5.15, 3.7, 0.9, align=PP_ALIGN.CENTER)
    rich(tf, [("10.7% ", True, BLUE),
              ("of the variation in the score is due to the group",
               False, TEXT)], size=14, first=True, space_after=0)
    math(s, r"\eta^2 = \frac{SS_{\mathrm{between\ groups}}}"
            r"{SS_{\mathrm{total}}} = 0.107", 4.75, 1.5, size=22)
    tf = textbox(s, 4.75, 2.55, 7.8, 0.3)
    para(tf, "η² ON EACH RUN (one-way ANOVA), AGAINST THE SMALLEST WE COULD "
             "DETECT", size=11, bold=True, colour=MUTED, first=True,
         space_after=0)
    scale_bars(s, [("first run", 0.1068, "0.107", BLUE),
                   ("second run", 0.0888, "0.089", MIDBLUE),
                   ("detectable", 0.0268, "0.027", MUTED)],
               top=3.0, maximum=0.12, left=6.35, length=4.6, label_w=1.45,
               step=0.55, size=14)
    stat(s, 4.75, 4.75, 3.8, 1.3, "F = 15.779", "between-group spread ÷ "
         "within-group spread", colour=MIDBLUE, size=24, label_size=12)
    stat(s, 8.78, 4.75, 3.8, 1.3, "p = 1.0 × 10⁻⁹", "probability of F this "
         "large if all 4 groups were equal", colour=GREEN, size=24,
         label_size=12)
    tf = textbox(s, M, 6.15, 6.0, 0.35)
    para(tf, "SS = sum of squares. The other 89.3% is article to article.",
         size=12, colour=MUTED, first=True, space_after=0)
    next_line(s, "which group drives it?")
    done(s, 6)

    # 8 result 2
    s = page(7)
    header(s, "The results", "Result 2: fear drives it, d = 0.94; outrage "
                             "lifts both sides equally", "chart-bar")
    math(s, r"d = \frac{\bar{X}_{\mathrm{group}} - \bar{X}_{\mathrm{neutral}}}"
            r"{s}", M, 1.45, size=20)
    tf = textbox(s, 3.9, 1.55, 8.5, 0.5)
    para(tf, "d = distance from neutral in standard deviations (s). "
             "Reference marks 0.2, 0.5, 0.8.", size=12, colour=MUTED,
         first=True, space_after=0)
    left, length = 2.9, 6.6
    for mark in (0.2, 0.5, 0.8):
        xm = left + length * mark
        bar(s, xm, 2.3, 0.015, 1.7, LINE)
        tf = textbox(s, xm - 0.6, 2.08, 1.2, 0.25, align=PP_ALIGN.CENTER)
        para(tf, "%.1f" % mark, size=11, colour=MUTED, first=True,
             space_after=0)
    scale_bars(s, [("fear-activating", 0.939, "0.94", EMBER),
                   ("reward hook", 0.319, "0.32", MIDBLUE),
                   ("high outrage", 0.030, "0.03", GOLD)],
               top=2.45, maximum=1.0, left=left, length=length, label_w=2.0,
               step=0.55)
    for x, name, emo, rea, verdict, colour in [
            (M, "FEAR", 0.613, -0.098, "score moves:\nd = 0.94", EMBER),
            (6.83, "OUTRAGE", 0.507, 0.491, "score flat:\nd = 0.03", GOLD)]:
        card(s, x, 4.2, 5.75, 2.0, line=colour)
        tf = textbox(s, x + 3.0, 4.35, 2.6, 0.3)
        para(tf, name + ", d per region", size=12, bold=True, colour=colour,
             first=True, space_after=0)
        base = 5.55
        bar(s, x + 0.3, base, 2.4, 0.02, MUTED)
        for j, (val, c, lab) in enumerate([(emo, EMBER, "emotional"),
                                           (rea, BLUE, "deliberate")]):
            bx = x + 0.5 + j * 1.05
            hgt = abs(val) * 1.6
            bar(s, bx, base - hgt if val > 0 else base, 0.7, max(hgt, 0.03), c)
            tf = textbox(s, bx - 0.2, base - hgt - 0.3 if val > 0 else base
                         - 0.3, 1.1, 0.3, align=PP_ALIGN.CENTER)
            para(tf, "%+.2f" % val, size=12, bold=True, colour=c, first=True,
                 space_after=0)
            tf = textbox(s, bx - 0.2, base + 0.2, 1.1, 0.3,
                         align=PP_ALIGN.CENTER)
            para(tf, lab, size=11, colour=MUTED, first=True, space_after=0)
        tf = textbox(s, x + 3.0, 4.75, 2.6, 1.2, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, verdict, size=17, bold=True, colour=colour, first=True,
             space_after=0)
    next_line(s, "can we trust the scores?")
    done(s, 7)

    # 9 result 3
    s = page(8)
    header(s, "The results", "Result 3: scores repeat for groups, not single "
                             "articles", "arrows-clockwise")
    card(s, M, 1.5, 5.75, 3.4, line=BLUE)
    tf = textbox(s, M, 1.65, 5.75, 1.2, align=PP_ALIGN.CENTER)
    para(tf, "0.8725", size=54, bold=True, colour=BLUE, first=True,
         space_after=0)
    para(tf, "ICC: agreement between two runs", size=14, colour=TEXT,
         space_after=0)
    bar(s, M + 0.4, 3.5, 4.95, 0.3, LINE)
    bar(s, M + 0.4, 3.5, 4.95 * 0.8725, 0.3, BLUE)
    for v, label in [(0, "0 = none"), (1, "1 = identical")]:
        tf = textbox(s, M + 0.4 + 4.95 * v - 0.2 - 1.0 * v, 3.9, 1.4, 0.3,
                     align=PP_ALIGN.CENTER)
        para(tf, label, size=12, colour=MUTED, first=True, space_after=0)
    card(s, 6.83, 1.5, 5.75, 3.4, line=EMBER)
    tf = textbox(s, 6.83, 1.65, 5.75, 1.2, align=PP_ALIGN.CENTER)
    para(tf, "51 of 400", size=54, bold=True, colour=EMBER, first=True,
         space_after=0)
    para(tf, "changed sign between runs (12.8%)", size=14, colour=TEXT,
         space_after=0)
    lo, hi, ax0, axw, ay = 40, 70, 7.2, 5.0, 3.65

    def at(v):
        return ax0 + (v - lo) / (hi - lo) * axw

    bar(s, ax0, ay, axw, 0.02, MUTED)
    bar(s, at(44), ay - 0.15, at(67) - at(44), 0.32, MIDBLUE, 0.4)
    bar(s, at(55.3) - 0.02, ay - 0.25, 0.04, 0.52, BLUE)
    plain(s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(at(51) - 0.12),
                             Inches(ay - 0.11), Inches(0.24), Inches(0.24)),
          EMBER)
    tf = textbox(s, 6.83, 4.1, 5.75, 0.6, align=PP_ALIGN.CENTER)
    rich(tf, [("● observed 51   ", True, EMBER),
              ("▌noise predicts 55 (95%: 44 to 67)", True, BLUE)], size=12,
         first=True, space_after=0)
    banner(s, 5.2, "Meaning:", "flipped articles average |X| = 0.003 against "
           "0.022 for the rest, so we report group averages only.", size=15)
    next_line(s, "does the score detect manipulation?")
    done(s, 8)

    # 10 honest check
    s = page(9)
    header(s, "The results", "Honest check: word counting beats our score, "
                             "0.98 vs 0.63", "warning")
    tf = textbox(s, M, 1.45, W - 2 * M, 0.35)
    para(tf, "AUC = chance a random manipulative article scores above a "
             "random neutral one", size=13, colour=MUTED, first=True,
         space_after=0)
    lo, hi, ax0, axw, ay = 0.5, 1.0, 1.3, 10.7, 2.75

    def au(v):
        return ax0 + (v - lo) / (hi - lo) * axw

    bar(s, ax0, ay, axw, 0.06, LINE)
    for v, label in [(0.5, "0.5 = coin toss"), (1.0, "1.0 = perfect")]:
        tf = textbox(s, au(v) - 0.9, ay + 0.2, 1.8, 0.3, align=PP_ALIGN.CENTER)
        para(tf, label, size=12, colour=MUTED, first=True, space_after=0)
    for v, label, colour, above in [(0.5392, "sentiment 0.54", MUTED, False),
                                    (0.6274, "our score 0.63", BLUE, True),
                                    (0.9758, "word counting 0.98", EMBER, True)]:
        plain(s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(au(v) - 0.16),
                                 Inches(ay - 0.13), Inches(0.32), Inches(0.32)),
              colour)
        tf = textbox(s, au(v) - 1.3, ay - 0.62 if above else ay + 0.5, 2.6, 0.4,
                     align=PP_ALIGN.CENTER)
        para(tf, label, size=16, bold=True, colour=colour, first=True,
             space_after=0)
    tf = textbox(s, M, 3.65, 5.5, 0.3)
    para(tf, "WHY: EACH GROUP CAME FROM ONE SOURCE", size=11, bold=True,
         colour=MUTED, first=True, space_after=0)
    for i, (grp, src) in enumerate([("fear", "ISOT fake"),
                                    ("outrage", "SemEval"),
                                    ("reward hook", "Webis"),
                                    ("neutral", "PubMed + ISOT")]):
        y = 4.0 + i * 0.5
        chip(s, M, y, 2.1, 0.42, grp, size=13, align=PP_ALIGN.CENTER)
        arrow(s, 2.95, y + 0.14, 0.5, 0.14, EMBER)
        chip(s, 3.55, y, 2.5, 0.42, src, size=13, align=PP_ALIGN.CENTER)
    stat(s, 6.6, 3.7, 2.8, 1.85, "66%", "of articles: words identify\n"
         "the source (chance 25%)", colour=EMBER, size=36)
    card(s, 9.6, 3.7, 2.98, 1.85, line=BLUE)
    tf = textbox(s, 9.75, 3.8, 2.7, 1.7, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "So our score is a measurement for the physics, not a detector.",
         size=15, bold=True, colour=BLUE, first=True, space_after=0)
    next_line(s, "does the brain model match real brains?")
    done(s, 9)

    # 11 real brains
    s = page(10)
    header(s, "The results", "Does the brain model match real brains? "
                             "r = +0.028, p = 0.048", "target")
    tf = textbox(s, M, 1.45, W - 2 * M, 0.35)
    para(tf, "Public brain scans of 4 people watching Friends (Algonauts "
             "2025).  r = correlation: 0 = no match, 1 = perfect.", size=13,
         colour=MUTED, first=True, space_after=0)
    scale_bars(s, [("people vs each other,\nfull episode", 0.1517, "0.152",
                    MUTED),
                   ("ceiling,\n2-minute clip", 0.0959, "0.096", BLUE),
                   ("our model,\n2-minute clip", 0.0283, "0.028", GREEN)],
               top=2.1, maximum=0.17, left=3.6, length=6.8, label_w=2.6,
               step=0.8, size=14)
    stat(s, M, 4.65, 3.75, 1.35, "30%", "of the clip ceiling\n"
         "(0.028 ÷ 0.096)", colour=GREEN, size=30)
    stat(s, 4.79, 4.65, 3.75, 1.35, "p = 0.048", "against 61 time-shifted "
         "copies;\ncut-off 0.05", colour=GOLD, size=30)
    stat(s, 8.83, 4.65, 3.75, 1.35, "3 bugs", "found and fixed along the "
         "way", colour=EMBER, size=30)
    banner(s, 6.1, "Meaning:", "the sign is positive (+0.028). The "
           "full-episode run is next, in Paper 3.", h=0.5, size=15)
    next_line(s, "the physics")
    done(s, 10)

    # 12 the minimum push
    s = page(11)
    header(s, "The physics", "The minimum push media would need: "
                             "α ≥ 4.29", "ruler")
    card(s, M, 1.45, 5.7, 4.75, line=BLUE)
    math(s, r"h = \alpha X \;\Rightarrow\; \alpha \geq "
            r"\frac{h_c(\beta J)}{\Delta X}", M, 1.65, size=28, centre_w=5.7)
    for i, (sym, name, text) in enumerate([
            ("h", "external field", "media's push on everyone"),
            ("X", "observable", "our score for an article"),
            (r"\alpha", "coupling constant", "push per unit of score"),
            (r"h_c", "critical field", "smallest push that flips the "
                                       "majority"),
            (r"\beta J", "reduced coupling", "how strongly people copy "
                                             "each other"),
            (r"\Delta X", "spread", "highest − lowest score = 0.124")]):
        y = 2.72 + i * 0.57
        math(s, sym, M + 0.15, y + 0.06, size=19, colour=GOLD, centre_w=0.95)
        tf = textbox(s, M + 1.2, y, 4.4, 0.55, anchor=MSO_ANCHOR.MIDDLE)
        rich(tf, [(name + "  ", True, BLUE), ("= " + text, False, TEXT)],
             size=13, first=True, space_after=0)
    alpha_curve(s, 6.75, 1.4, 5.9, 3.7)
    tf = textbox(s, 6.75, 5.0, 5.9, 0.3, align=PP_ALIGN.CENTER)
    para(tf, "reduced coupling βJ  →", size=12, colour=MUTED, first=True,
         space_after=0)
    tf = textbox(s, 6.95, 1.35, 3, 0.3)
    para(tf, "minimum α", size=12, colour=MUTED, first=True, space_after=0)
    stat(s, 6.75, 5.35, 5.83, 0.85, "α ≥ 0.533 ÷ 0.124 = 4.29",
         "at βJ = 2, where the critical field is 0.533", colour=EMBER, size=22)
    next_line(s, "what it could be used for")
    done(s, 11)

    # 13 children's content
    s = page(12)
    header(s, "What's next", "Can it analyse any content, even what children "
                             "watch?", "target")
    for i, (head, colour, icon_name, items) in enumerate([
            ("CAN DO", GREEN, "check-circle",
             ["Score video, audio and text", "Compare groups of content",
              "Rate content, never the viewer"]),
            ("NOT YET", EMBER, "x-circle",
             ["Model learned from adult brains", "Only news text was tested",
              "1 in 8 single scores flip"]),
            ("WOULD TAKE", BLUE, "flag-banner",
             ["Child brain data, with ethics approval",
              "A children's video test set",
              "Safeguards against targeting children"])]):
        x = M + i * (3.75 + 0.29)
        card(s, x, 1.5, 3.75, 3.85, line=colour)
        icon(s, icon_name, x + 3.75 / 2 - 0.3, 1.7, h=0.6)
        tf = textbox(s, x, 2.4, 3.75, 0.4, align=PP_ALIGN.CENTER)
        para(tf, head, size=16, bold=True, colour=colour, first=True,
             space_after=0)
        for j, text in enumerate(items):
            chip(s, x + 0.2, 2.95 + j * 0.78, 3.35, 0.66, text, size=14,
                 align=PP_ALIGN.CENTER)
    banner(s, 5.6, "Short answer:", "technically possible, but not yet valid "
           "for children.", h=0.6)
    next_line(s, "the limits, and what fixes each")
    done(s, 12)

    # 14 limits
    s = page(13)
    header(s, "What's next", "Each limit points to the next study", "x-circle")
    for x, label, colour in [(M, "LIMIT", EMBER), (6.2, "NEXT STUDY", GREEN)]:
        tf = textbox(s, x, 1.45, 4, 0.3)
        para(tf, label, size=11, bold=True, colour=colour, first=True,
             space_after=0)
    for i, (limit, nxt) in enumerate([
            ("Predictions, not brain scans", "Check against more real brain "
                                             "data"),
            ("Group mixed up with source", "New set: every source has every "
                                           "type"),
            ("Brain surface only", "A model that predicts deep areas"),
            ("α not measured", "Link scores to real opinion data (polls)"),
            ("Brain check is 2 minutes", "The full episode")]):
        y = 1.8 + i * 0.86
        chip(s, M, y, 4.85, 0.72, limit, line=EMBER, bold=True)
        arrow(s, 5.65, y + 0.29, 0.45, 0.16, GREEN)
        chip(s, 6.2, y, 6.38, 0.72, nxt, line=GREEN)
    next_line(s, "where the work goes")
    done(s, 13)

    # 15 taking it further
    s = page(14)
    header(s, "What's next", "Taking it further: where this research can go",
           "flag-banner")
    for i, (head, colour, icon_name, items) in enumerate([
            ("PUBLISH", BLUE, "books",
             ["Paper 1 → Physica A", "Paper 2 → Physica A",
              "Paper 3 → Imaging Neuroscience", "Preprints first (arXiv)"]),
            ("EXTEND THE SCIENCE", GREEN, "atom",
             ["Kenya: 2024 Finance Bill (funding)", "Video and audio content",
              "Measure α with opinion data", "Children's media, ethically"]),
            ("OPPORTUNITIES", GOLD, "target",
             ["Postgraduate research topic", "Neuroscience + media "
              "collaborations", "Open tool, already online",
              "Group-level media audits"])]):
        x = M + i * (3.75 + 0.29)
        card(s, x, 1.5, 3.75, 4.55, line=colour)
        icon(s, icon_name, x + 0.25, 1.7, h=0.45)
        tf = textbox(s, x + 0.85, 1.72, 2.8, 0.45)
        para(tf, head, size=15, bold=True, colour=colour, first=True,
             space_after=0)
        for j, text in enumerate(items):
            chip(s, x + 0.2, 2.45 + j * 0.88, 3.35, 0.74, text, size=14)
    next_line(s, "conclusions")
    done(s, 14)

    # 16 conclusions
    s = page(15)
    header(s, "What's next", "Conclusions", "check-circle")
    for i, (value, colour, claim, meaning) in enumerate([
            ("0.8725", BLUE, "A field observable can be measured",
             "ICC: agreement between two runs"),
            ("0.107", GREEN, "It separates the four groups",
             "η²: 10.7% of variation due to group; source mixed in"),
            ("α ≥ 4.29", EMBER, "Media strength has a minimum",
             "at βJ = 2, within the mean-field model"),
            ("r = 0.028", GOLD, "The brain model matches real brains",
             "positive; 30% of the 0.096 ceiling")]):
        x = M + (i % 2) * 6.08
        y = 1.5 + (i // 2) * 1.85
        card(s, x, y, 5.75, 1.65, line=colour)
        tf = textbox(s, x + 0.2, y + 0.1, 0.5, 0.5)
        para(tf, str(i + 1), size=20, bold=True, colour=colour, first=True,
             space_after=0)
        tf = textbox(s, x + 0.7, y + 0.12, 4.9, 1.45)
        para(tf, value, size=30, bold=True, colour=colour, first=True,
             space_after=2)
        para(tf, claim, size=15, bold=True, colour=TEXT, space_after=0)
        para(tf, meaning, size=12, colour=MUTED, space_after=0)
    banner(s, 5.35, "Three papers follow.", "Paper 1 and Paper 2 → Physica A,  "
           "Paper 3 → Imaging Neuroscience.", colour=BLUE, h=0.55, size=14)
    tf = textbox(s, M, 6.05, W - 2 * M, 0.5, align=PP_ALIGN.CENTER)
    para(tf, "Thank you.  Questions welcome.", size=22, bold=True,
         colour=GOLD, first=True, align=PP_ALIGN.CENTER, space_after=0)
    done(s, 15)

    name = "Monarch_Defence_Deck_Script.pptx" if with_notes \
        else "Monarch_Defence_Deck.pptx"
    out = os.path.join(HERE, name)
    prs.save(out)
    return out, len(prs.slides)


def write_script():
    total = sum(entry[2] for entry in SLIDES)
    blocks = ["MONARCH DEFENCE SCRIPT  (about %g min)" % total]
    for i, entry in enumerate(SLIDES):
        blocks.append("=" * 72 + "\nSLIDE %d  %s\n" % (i + 1, entry[0])
                      + notes_for(i))
    out = os.path.join(HERE, "Monarch_Defence_Script.txt")
    with open(out, "w", encoding="utf-8") as handle:
        handle.write("\n\n".join(blocks) + "\n")
    return out


if __name__ == "__main__":
    for path, count in (build_deck(False), build_deck(True)):
        print("wrote", path, count, "slides")
    print("wrote", write_script())
